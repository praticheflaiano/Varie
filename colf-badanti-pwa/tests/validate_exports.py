import json
from pathlib import Path
from zipfile import ZipFile
import xml.etree.ElementTree as ET
from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader
root=Path("pwa-test-artifacts")
expected=json.loads((root/"expected.json").read_text())
for name in ("sample.docx","sample.xlsx"):
    with ZipFile(root/name) as z:
        assert z.testzip() is None
        for part in z.namelist():
            if part.endswith((".xml",".rels")):
                ET.fromstring(z.read(part))
doc=Document(root/"sample.docx")
assert len(doc.tables)>=8
assert any("Prova & Figli <Roma>" in c.text for t in doc.tables for row in t.rows for c in row.cells)
book=load_workbook(root/"sample.xlsx",data_only=True)
assert abs(book["Calcoli"]["B18"].value-expected["cost"])<0.001
assert abs(book["Calcoli"]["B19"].value-expected["net"])<0.001
assert abs(book["Calcoli"]["B24"].value-expected["annual"])<0.001
formulas=load_workbook(root/"sample.xlsx",data_only=False)
assert formulas["Calcoli"]["B18"].value.startswith("=")
pdf=PdfReader(root/"sample.pdf")
assert len(pdf.pages)>=2
text="\n".join(p.extract_text() or "" for p in pdf.pages)
assert "FLAIANO" in text
assert "Prova & Figli <Roma>" in text
print("DOCX, XLSX e PDF verificati con parser indipendenti.")

for name in ("contributi-trimestre.docx","contributi-trimestre.xlsx"):
    with ZipFile(root/name) as z:
        assert z.testzip() is None
        for part in z.namelist():
            if part.endswith((".xml",".rels")):
                ET.fromstring(z.read(part))
quarter_doc=Document(root/"contributi-trimestre.docx")
assert len(quarter_doc.tables)>=6
assert any("Famiglia & <prova>" in c.text for t in quarter_doc.tables for row in t.rows for c in row.cells)
quarter_book=load_workbook(root/"contributi-trimestre.xlsx",data_only=True)
assert quarter_book["Settimane"].max_row==14
assert quarter_book["Calcoli"]["B2"].value==325
assert quarter_book["Calcoli"]["B5"].value==403
assert abs(quarter_book["Calcoli"]["B15"].value-422.5)<0.001
assert quarter_book["Riconciliazione"]["B4"].value in (None,"")
quarter_formulas=load_workbook(root/"contributi-trimestre.xlsx",data_only=False)
assert quarter_formulas["Calcoli"]["B2"].value=="=ROUNDUP(SUM(Settimane!C2:C14),0)"
quarter_pdf=PdfReader(root/"contributi-trimestre.pdf")
quarter_text="\n".join(p.extract_text() or "" for p in quarter_pdf.pages)
assert "422,50" in quarter_text
assert "pagoPA" in quarter_text
print("Prospetti trimestrali DOCX, XLSX e PDF verificati con parser indipendenti.")
